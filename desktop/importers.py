from pathlib import Path
import csv

def read_txt(path):
    return [x.strip() for x in Path(path).read_text(encoding="utf-8-sig").splitlines() if x.strip()]

def read_csv_file(path):
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            if row and str(row[0]).strip():
                rows.append(str(row[0]).strip())
    return rows

def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return [str(row[0].value).strip() for row in ws.iter_rows() if row and row[0].value]

def import_prompts(path):
    ext = Path(path).suffix.lower()
    if ext == ".txt":
        return read_txt(path)
    if ext == ".csv":
        return read_csv_file(path)
    if ext in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError("Chỉ hỗ trợ TXT, CSV, XLSX.")
